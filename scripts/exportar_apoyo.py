#!/usr/bin/env python3
"""Lee las hojas de apoyo de Logistica_Viajes_DSY.xlsm y genera data/apoyo.json.
Uso: python scripts/exportar_apoyo.py Logistica_Viajes_DSY.xlsm   (requiere: pip install openpyxl)"""
import sys, json, datetime as dt
from openpyxl import load_workbook

COLS = {
 "tbAreas": None, "tbSaludo": None, "tbProductos": None, "tbCoordinador": None, "tbMuelles": None,
 "tbHoraOrden": None, "tbRespo": None, "tbDuenos": None,
 "tbPlaca": ["Placas", "Proveedor", "Sellos", "Capacidad", "C1", "C2", "C3", "C4", "C5", "C6"],
 "tbChofer": ["Chofer", "NombrePila", "Cedula", "Celular", "PLACAS"],
}
wb = load_workbook(sys.argv[1], data_only=True)
out = {}
for ws in wb.worksheets:
    for t in ws.tables.values():
        if t.name in COLS:
            cells = ws[t.ref]
            hdr = [str(c.value) for c in cells[0]]
            rows = []
            for r in cells[1:]:
                row = {}
                for h, c in zip(hdr, r):
                    if COLS[t.name] and h not in COLS[t.name]:
                        continue
                    v = c.value
                    if isinstance(v, dt.time): v = v.strftime("%H:%M")
                    elif isinstance(v, dt.datetime): v = v.isoformat()
                    row[h] = v
                if any(x not in (None, "") for x in row.values()):
                    rows.append(row)
            out[t.name] = rows
# listas para autocompletar (desde 2025) y texto fijo de Recordar
desde = dt.datetime(2025, 1, 1)
cl, se, de = set(), set(), set()
for r in wb["Programacion"].iter_rows(min_row=4, values_only=True):
    if isinstance(r[0], dt.datetime) and r[0] >= desde:
        cl.add(r[2]); se.add(r[3]); de.add(r[9])
f = lambda s: sorted(x for x in s if x)
out["listas"] = {"clientes": f(cl), "sectores": f(se), "destinos": f(de)}
out["recordatorio"] = wb["data"]["R1"].value or ""
json.dump(out, open("data/apoyo.json", "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
print({k: (len(v) if isinstance(v, list) else "ok") for k, v in out.items()})
